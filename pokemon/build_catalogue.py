#!/usr/bin/env python3
# Run from the repo root (the folder containing pokemon/).
#
# Windows (PowerShell / VS Code terminal):
#   pokemon\.venv\Scripts\python.exe pokemon\build_catalogue.py
#   pokemon\.venv\Scripts\python.exe pokemon\build_catalogue.py --refresh
# macOS:
#   pokemon/.venv/bin/python pokemon/build_catalogue.py
#
# First time on a machine only, to create the venv:
#   Windows:  py -3 -m venv pokemon\.venv
#             pokemon\.venv\Scripts\python.exe -m pip install -r pokemon\requirements.txt
#   macOS:    python3 -m venv pokemon/.venv
#             pokemon/.venv/bin/pip install -r pokemon/requirements.txt
"""Build the Pokémon TCG catalogue data for the Jekyll site.

Reads pokemon/source/pokemon_bulk_catalogue.xlsx, fetches the full card list of
every set in it from TCGCSV (TCGplayer mirror) and set metadata from TCGdex,
and writes:

  _data/pokemon/sets.json      one summary entry per owned set
  _data/pokemon/<slug>.json    every card + variant of that set, with qty
  pokemon-tcg/<slug>.md        stub page per set (GitHub Pages can't generate
                               pages from data)

Usage:  python pokemon/build_catalogue.py [--refresh]
"""

import argparse
import difflib
import json
import re
import sys
import time
import unicodedata
import warnings
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor
from datetime import date
from pathlib import Path

warnings.filterwarnings("ignore", message=".*OpenSSL.*")               # macOS system Python + urllib3 v2
warnings.filterwarnings("ignore", message=".*Data Validation extension.*")  # xlsx dropdowns

import openpyxl  # noqa: E402
import requests  # noqa: E402
import yaml  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
SOURCE_DIR = ROOT / "pokemon" / "source"
XLSX = SOURCE_DIR / "pokemon_bulk_catalogue.xlsx"
OVERRIDES = SOURCE_DIR / "set_overrides.yml"
NON_COLLECTION_SHEETS = {"How to use", "Lists"}
CACHE_DIR = ROOT / "pokemon" / ".cache"
DATA_DIR = ROOT / "_data" / "pokemon"
PAGES_DIR = ROOT / "pokemon-tcg"
VERBOSE = True  # progress lines like "checking N image(s)"; master set scripts turn them off

TCGCSV = "https://tcgcsv.com/tcgplayer/3"
TCGDEX = "https://api.tcgdex.net/v2/en"
PTCG = "https://api.pokemontcg.io/v2"

# Order here is also the display order of variants for the same card number.
VARIANTS = {
    "normal": "Normal",
    "holo": "Holo",
    "reverse": "Reverse",
    "1st-normal": "1st Edition",        # vintage sets: 1st Edition vs Unlimited (= normal / holo)
    "1st-holo": "1st Edition Holo",
    "pokeball": "Poké Ball",
    "masterball": "Master Ball",
    "friendball": "Friend Ball",
    "loveball": "Love Ball",
    "quickball": "Quick Ball",
    "duskball": "Dusk Ball",
    "energysymbol": "Energy Symbol",
    "rocket": "Team Rocket",
    # special prints outside the booster set (TCGplayer's exclusives / promo groups)
    "cosmos": "Cosmos Holo",
    "altholo": "Alternate Holo",
    "stamp": "Stamped",
    "prerelease": "Prerelease",
    "staff": "Staff",
    "promo": "Promo",
}

# Groups holding a set's special prints under the set's own numbers (031/162, ...)
EXCLUSIVE_GROUPS = ("blister exclusives", "miscellaneous cards and products")

# TCGplayer price subTypeName of the plain product -> variant
SUBTYPES = {
    "Normal": "normal",
    "Holofoil": "holo",
    "Reverse Holofoil": "reverse",
    "Unlimited": "normal",
    "Unlimited Normal": "normal",
    "Unlimited Holofoil": "holo",
    "1st Edition": "1st-normal",
    "1st Edition Normal": "1st-normal",
    "1st Edition Holofoil": "1st-holo",
}

# Bracketed label in a TCGplayer product name -> patterned-reverse variant
PATTERNS = [
    (re.compile(r"^pok[eé] ?ball( pattern)?$", re.I), "pokeball"),
    (re.compile(r"^master ?ball( pattern)?$", re.I), "masterball"),
    (re.compile(r"^friend ?ball( pattern)?$", re.I), "friendball"),
    (re.compile(r"^love ?ball( pattern)?$", re.I), "loveball"),
    (re.compile(r"^quick ?ball( pattern)?$", re.I), "quickball"),
    (re.compile(r"^dusk ?ball( pattern)?$", re.I), "duskball"),
    (re.compile(r"^energy symbol( pattern)?$", re.I), "energysymbol"),
    (re.compile(r"^(team )?rocket( pattern)?$", re.I), "rocket"),
]

session = requests.Session()
session.headers["User-Agent"] = "camattelaer.github.io pokemon catalogue builder"


# ---------------------------------------------------------------- helpers

def get_json(url, cache_name, refresh=False):
    """GET url as JSON, cached in pokemon/.cache/<cache_name>.json."""
    path = CACHE_DIR / f"{cache_name}.json"
    if path.exists() and not refresh:
        return json.loads(path.read_text())
    for attempt in range(3):
        try:
            resp = session.get(url, timeout=30)
            resp.raise_for_status()
            data = resp.json()
            break
        except requests.RequestException as exc:
            if attempt == 2:
                raise
            print(f"  retrying {url} ({exc})", file=sys.stderr)
            time.sleep(2 * (attempt + 1))
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data))
    return data


def fold(text):
    """Lowercase, strip accents and punctuation, collapse whitespace."""
    text = unicodedata.normalize("NFKD", str(text)).encode("ascii", "ignore").decode()
    text = text.lower().replace("&", " and ")
    return " ".join(re.sub(r"[^a-z0-9]+", " ", text).split())


def slugify(text):
    return fold(text).replace(" ", "-")


def number_key(raw):
    """Canonical number for matching: '025/217' -> '25', 'TG05/TG30' -> 'TG5'."""
    raw = str(raw).strip().split("/")[0].strip().upper().replace(" ", "")  # 'MEP 084' -> 'MEP084'
    if raw.endswith(".0") and raw[:-2].isdigit():  # xlsx numeric cells
        raw = raw[:-2]
    m = re.fullmatch(r"([A-Z-]*)0*(\d+)([A-Z]*)", raw)
    if m:
        return f"{m.group(1)}{m.group(2)}{m.group(3)}"
    return raw


def full_key(raw):
    """number_key plus the printed set total when there is one: '106/105' -> '106/105',
    '025' -> '25'. Needed for sets numbered as the original prints (Classic Collection),
    where several cards share a number."""
    text = str(raw).strip()
    total = text.split("/")[1].strip() if "/" in text else ""
    return f"{number_key(text)}/{number_key(total)}" if total else number_key(text)


def key_of_number(number, reprint):
    """Card key of a printed number as build_card_list keys it."""
    return full_key(number) if reprint else number_key(number)


def number_sort(key):
    """Plain numbers first in numeric order, then prefixed ones (TG, GG, SV...);
    full keys ('106/105') by number, then set total."""
    key, _, total = key.partition("/")
    m = re.fullmatch(r"([A-Z-]*)(\d+)([A-Z]*)", key)
    if not m:
        return (2, key, 0, "", total)
    prefix, digits, suffix = m.groups()
    return (0 if not prefix else 1, prefix, int(digits), suffix, number_sort(total) if total else ())


def strip_group_prefix(name):
    """'SV: Prismatic Evolutions' -> 'Prismatic Evolutions', 'SM - Lost Thunder' -> 'Lost Thunder'."""
    return re.sub(r"^[A-Za-z0-9&.]+(\s*:\s*|\s+-\s+)", "", name).strip()


def product_number(product):
    for field in product.get("extendedData", []):
        if field["name"] == "Number":
            return field["value"]
    return None


def clean_card_name(name):
    name = re.sub(r"\s*\([^)]*\)", "", name)
    name = re.sub(r"\s+-\s+[A-Z]*\d+[A-Z]*(/[A-Z]*\d+)?$", "", name)
    return name.strip()


# ---------------------------------------------------------------- inputs

def read_catalogue():
    """One sheet per set: set_name in B1, then number / variant / qty from row 3 down.
    Every sheet except 'How to use' and 'Lists' is read; empty ones (Template) are skipped."""
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    owned = defaultdict(int)  # (set_name, number_key, variant) -> qty
    rows_of = defaultdict(list)  # same key -> "'Sheet' row N" locations, for the report
    problems = []
    sheets_of = defaultdict(list)  # set_name -> sheet titles, to flag a set entered twice

    for ws in wb.worksheets:
        if ws.title in NON_COLLECTION_SHEETS:
            continue
        rows = ws.iter_rows(values_only=True)
        top = next(rows, ())
        set_name = str(top[1] if len(top) > 1 and top[1] is not None else "").strip()
        next(rows, None)
        header = [str(h).strip().lower() if h is not None else "" for h in next(rows, ())]
        if not all(name in header for name in ("number", "variant", "qty")):
            problems.append(f"sheet '{ws.title}': row 3 must be the headers number / variant / qty")
            continue
        col = {name: header.index(name) for name in ("number", "variant", "qty")}

        found_rows = False
        for rownum, row in enumerate(rows, start=4):
            number, variant, qty = (row[col[c]] if col[c] < len(row) else None
                                    for c in ("number", "variant", "qty"))
            if all(v in (None, "") for v in (number, variant, qty)):
                continue
            found_rows = True
            where = f"'{ws.title}' row {rownum}"
            variant = str(variant or "").strip().lower()
            if not set_name:
                continue  # reported once per sheet below
            if number in (None, ""):
                problems.append(f"{where}: missing number")
                continue
            if variant not in VARIANTS:
                problems.append(f"{where}: {set_name} #{number}: unknown variant '{variant}'")
                continue
            try:
                qty = int(qty) if qty not in (None, "") else 1
            except (TypeError, ValueError):
                problems.append(f"{where}: {set_name} #{number}: qty '{qty}' is not a number")
                continue
            if qty <= 0:
                continue
            owned[(set_name, full_key(number), variant)] += qty
            rows_of[(set_name, full_key(number), variant)].append(where)
        if found_rows and not set_name:
            problems.append(f"sheet '{ws.title}': no set_name in cell B1; its rows are skipped")
        if set_name:
            sheets_of[set_name].append(ws.title)

    for set_name, titles in sheets_of.items():
        if len(titles) > 1:
            print(f"warning: {set_name} is on several sheets ({', '.join(titles)}); quantities are added up",
                  file=sys.stderr)

    # Lists sheet: fallback series / release date per set name
    lists = {}
    if "Lists" in wb.sheetnames:
        for row in wb["Lists"].iter_rows(min_row=2, values_only=True):
            if row and row[0]:
                lists[str(row[0]).strip()] = {
                    "series": row[1] if len(row) > 1 else None,
                    "release_date": str(row[2])[:10] if len(row) > 2 and row[2] else None,
                }
    wb.close()
    return owned, rows_of, lists, problems


def read_overrides():
    if not OVERRIDES.exists():
        return {}
    return yaml.safe_load(OVERRIDES.read_text()) or {}


# ---------------------------------------------------------------- matching

class Catalogues:
    def __init__(self, refresh):
        self.refresh = refresh
        self._ptcg_sets = None  # fetched on first missing logo
        self._load(refresh)

    def _load(self, refresh):
        self.tcgdex_sets = get_json(f"{TCGDEX}/sets", "tcgdex_sets", refresh)
        self.groups = get_json(f"{TCGCSV}/groups", "tcgcsv_groups", refresh)["results"]
        self.loaded_fresh = refresh

    def reload_once(self):
        """Cached lists may predate a new set: refetch them once on a miss."""
        if self.loaded_fresh:
            return False
        print("  (refreshing cached set lists)")
        self._load(True)
        return True

    def tcgdex_detail(self, set_id):
        return get_json(f"{TCGDEX}/sets/{set_id}", f"tcgdex_set_{set_id}", self.refresh)

    def fallback_logo(self, name, release_date):
        """TCGdex lacks some logos (e.g. Temporal Forces); try pokemontcg.io by set name."""
        if self._ptcg_sets is None:
            try:
                self._ptcg_sets = get_json(f"{PTCG}/sets?pageSize=250&select=name,releaseDate,images",
                                           "ptcg_sets", self.refresh)["data"]
            except requests.RequestException as exc:
                print(f"  warning: pokemontcg.io unavailable, no fallback logos ({exc})", file=sys.stderr)
                self._ptcg_sets = []
        matches = [s for s in self._ptcg_sets if fold(s["name"]) == fold(name)]
        if len(matches) > 1 and release_date:
            matches = [s for s in matches if s.get("releaseDate", "").replace("/", "-") == release_date] or matches
        return matches[0]["images"].get("logo") if matches else None

    def match_tcgdex(self, name, override):
        if override:
            return next((s for s in self.tcgdex_sets if s["id"] == str(override)), None), "override"
        exact = [s for s in self.tcgdex_sets if fold(s["name"]) == fold(name)]
        if len(exact) == 1:
            return exact[0], "exact"
        # xlsx set names come from TCGdex, so only accept a very close fuzzy match
        names = {fold(s["name"]): s for s in self.tcgdex_sets}
        close = difflib.get_close_matches(fold(name), names, n=1, cutoff=0.9)
        return (names[close[0]], "fuzzy") if close else (None, None)

    def match_group(self, name, release_date, override):
        if override:
            return next((g for g in self.groups if g["groupId"] == int(override)), None), "override"
        target = fold(name)
        exact = [g for g in self.groups if fold(strip_group_prefix(g["name"])) == target]
        if len(exact) == 1:
            return exact[0], "exact"

        scored = []
        for g in self.groups:
            cand = fold(strip_group_prefix(g["name"]))
            score = difflib.SequenceMatcher(None, target, cand).ratio()
            if target and re.search(rf"\b{re.escape(target)}\b", cand):
                score += 0.15
            if release_date and g.get("publishedOn"):
                try:
                    days = abs((date.fromisoformat(g["publishedOn"][:10])
                                - date.fromisoformat(release_date)).days)
                    if days <= 21:
                        score += 0.15
                except ValueError:
                    pass
            scored.append((score, g))
        scored.sort(key=lambda s: -s[0])
        if scored and scored[0][0] >= 0.85 and (len(scored) == 1 or scored[0][0] - scored[1][0] >= 0.05):
            return scored[0][1], "fuzzy"
        return None, None

    def subset_groups(self, group, override):
        """TCGplayer splits Trainer Gallery, Galarian Gallery, Shiny Vault and Radiant
        Collection subsets into their own groups; TCGdex keeps them in the main set."""
        if override is not None:
            ids = {int(i) for i in override}
            return [g for g in self.groups if g["groupId"] in ids]
        pattern = re.compile(re.escape(group["name"]) +
                             r"(\s+Trainer Gallery|:\s*(Galarian Gallery|Shiny Vault|Radiant Collection))$")
        return [g for g in self.groups if pattern.match(g["name"])]


# ---------------------------------------------------------------- set build

def match_by_name(cards, tcgdex_cards):
    """{(number, name): TCGdex card} for a set numbered as the original prints, whose
    TCGdex numbers (001, 002, ...) differ from the printed ones. Same-named cards
    (Darkrai & Cresselia LEGEND top and bottom) pair up in number order; 'Palkia' on
    TCGdex matches 'Palkia LV.X'."""
    pool = sorted(tcgdex_cards, key=lambda t: number_sort(number_key(t["localId"])))
    result = {}
    unique = sorted({(c["number"], c["name"]) for c in cards}, key=lambda nc: number_sort(full_key(nc[0])))
    for exact in (True, False):
        for number, name in unique:
            if (number, name) in result:
                continue
            want = fold(name)
            hit = next((t for t in pool if (fold(t["name"]) == want if exact else
                                             want.startswith(fold(t["name"]) + " "))), None)
            if hit:
                result[(number, name)] = hit
                pool.remove(hit)
    return result


def special_kind(labels):
    """Variant of a special print from its TCGplayer labels: staff, prerelease, cosmos,
    altholo (Mirage / Line / Sparkle Holo, ...), stamp, else promo (store and event promos:
    GameStop, Pokemon Center, Holiday Calendar, ...)."""
    text = " ".join(labels).lower()
    if "staff" in text:
        return "staff"
    if "prerelease" in text:
        return "prerelease"
    if "cosmos" in text or "cosmo holo" in text:
        return "cosmos"
    if "holo" in text:
        return "altholo"
    if "stamp" in text:
        return "stamp"
    return "promo"


def special_prints(set_name, official, cards, groups, refresh):
    """A set's prints outside its own TCGplayer group, keyed (key, variant):
    - Blister Exclusives / Miscellaneous: Cosmos Holos and stamps that carry the set's own
      number (020/084), matched on printed total + number + name
    - promo groups: promos stamped with the set's name ('Pitch Black Stamped') mark a number
      range in that group (MEP 082-085); everything in it belongs to the set (staff versions,
      the prerelease promo in between), plus a prerelease just outside it of a card in the set.
      Promos keep their own number with the group's prefix ('MEP 084')."""
    found, notes = {}, []
    names = {fold(c["name"]) for c in cards.values()}
    by_number = {(c["key"], fold(c["name"])) for c in cards.values()}

    def entry(prod, number, key, variant, label):
        return {
            "number": number, "key": key,
            "name": clean_card_name(re.sub(r"\s*\[[^\]]*\]", "", prod["name"])),  # drop '[Staff]'
            "rarity": next((f["value"] for f in prod.get("extendedData", []) if f["name"] == "Rarity"), None),
            "image": (prod.get("imageUrl") or "").replace("_200w.", "_400w.") if prod.get("imageCount") else None,
            "url": prod.get("url"), "product_id": prod["productId"], "variant": variant, "special": label,
        }

    def add(card):
        """Same card, same kind sold under several labels (GameStop / EB Games / European
        promo): one print, with all labels."""
        k = (card["key"], card["variant"])
        if k in found:
            if card["special"] not in found[k]["special"].split(" / "):
                found[k]["special"] += " / " + card["special"]
        else:
            found[k] = card

    def labels_of(prod):
        """'Bastiodon - 085 (Pitch Black Stamped) [Staff]' -> ['Pitch Black Stamped', 'Staff']"""
        pairs = re.findall(r"\(([^)]*)\)|\[([^\]]*)\]", prod["name"])
        return [(round_ or square).strip() for round_, square in pairs if (round_ or square).strip()]

    for group in groups:
        gname = fold(group["name"])
        promo = "promo" in gname
        if not promo and gname not in EXCLUSIVE_GROUPS:
            continue
        try:
            products = get_json(f"{TCGCSV}/{group['groupId']}/products", f"tcgcsv_{group['groupId']}_products",
                                refresh)["results"]
        except requests.RequestException:
            continue
        if not promo:
            for prod in products:
                raw = product_number(prod) or ""
                total = raw.split("/")[1].strip() if "/" in raw else ""
                if not (total.isdigit() and int(total) == official
                        and (number_key(raw), fold(clean_card_name(prod["name"]))) in by_number):
                    continue
                labels = labels_of(prod)
                kind = special_kind(labels)
                add(entry(prod, raw.split("/")[0].strip(), number_key(raw), kind,
                          ", ".join(labels) or group["name"]))
            continue

        stamped = [prod for prod in products if fold(set_name) in fold(" ".join(labels_of(prod)))]
        numbers = [int(number_key(product_number(p) or "x")) for p in stamped
                   if number_key(product_number(p) or "x").isdigit()]
        if not numbers:
            continue
        low, high = min(numbers), max(numbers)
        prefix = (group.get("abbreviation") or "PROMO").upper()
        for prod in products:
            raw = product_number(prod) or ""
            n = number_key(raw)
            if not n.isdigit():
                continue
            labels = labels_of(prod)
            near_prerelease = (low - 2 <= int(n) <= high + 2 and "prerelease" in " ".join(labels).lower()
                               and fold(clean_card_name(prod["name"])) in names)
            if not (low <= int(n) <= high or near_prerelease):
                continue
            kind = special_kind(labels) if labels else "holo"
            key = f"{prefix}{n}"
            add(entry(prod, f"{prefix} {raw.strip()}", key, kind, ", ".join(labels) or group["name"]))
    return found, notes


def check_images(urls, refresh):
    """url -> whether the CDN serves it. Cached in pokemon/.cache/image_status.json."""
    path = CACHE_DIR / "image_status.json"
    status = {} if refresh or not path.exists() else json.loads(path.read_text())

    def probe(url):
        try:
            resp = session.head(url, timeout=15, allow_redirects=True)
            ok = resp.ok and resp.headers.get("Content-Type", "").startswith("image/")
            return url, ok
        except requests.RequestException:
            return url, None  # network trouble: don't cache, assume fine for now

    todo = [u for u in urls if u not in status]
    if todo:
        if VERBOSE:
            print(f"  checking {len(todo)} image(s)...")
        with ThreadPoolExecutor(max_workers=16) as pool:
            for url, ok in pool.map(probe, todo):
                if ok is not None:
                    status[url] = ok
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(json.dumps(status))
    return {u: status.get(u, True) for u in urls}


def build_card_list(group_ids, refresh, deck_group_id=None, official=None):
    """Every card + variant in the given TCGplayer groups, keyed by (number_key, variant),
    plus the set's deck-exclusive prints (official = the set's official card count).
    Also returns whether the set uses original-print numbering (keys are then full_key)."""
    products, subtypes = [], defaultdict(set)
    for group_id in group_ids:
        products += get_json(f"{TCGCSV}/{group_id}/products", f"tcgcsv_{group_id}_products", refresh)["results"]
        for p in get_json(f"{TCGCSV}/{group_id}/prices", f"tcgcsv_{group_id}_prices", refresh)["results"]:
            subtypes[p["productId"]].add(p["subTypeName"])

    # Sort products into plain ones (variants come from price subtypes), ball/energy
    # pattern reverses, and other labels. Labels like "(Full Art)" or "(Secret)" in
    # older sets are separate cards with their own number; labels on a number that
    # already has a plain product ("(Prerelease)", stamps, ...) are skipped.
    plain, patterned, labelled = [], [], []
    for prod in products:
        raw_number = product_number(prod)
        if not raw_number:
            continue  # sealed product
        labels = [l.strip() for l in re.findall(r"\(([^)]*)\)", prod["name"])]
        variant = next((v for label in labels for rx, v in PATTERNS if rx.match(label)), None)
        entry = (prod, raw_number, labels)
        if variant:
            patterned.append(entry + (variant,))
        elif labels:
            labelled.append(entry)
        else:
            plain.append(entry)

    # Sets like a Classic Collection print each card with its original number and set
    # total (114/264, 138/202, ...), so numbers repeat: key those by the full number.
    totals = {n.split("/")[1].strip() for _, n, _ in plain + labelled if "/" in n}
    reprint = len(totals) > 3
    key_of = full_key if reprint else number_key

    plain_numbers = {key_of(n) for _, n, _ in plain}
    unrecognised = defaultdict(int)
    for prod, raw_number, labels in labelled:
        if key_of(raw_number) in plain_numbers:
            unrecognised[", ".join(labels)] += 1
        else:
            plain.append((prod, raw_number, labels))
            plain_numbers.add(key_of(raw_number))

    def card(prod, raw_number, variant):
        return {
            "number": raw_number.strip() if reprint else raw_number.split("/")[0].strip(),
            "key": key_of(raw_number),
            "name": clean_card_name(prod["name"]),
            "rarity": next((f["value"] for f in prod.get("extendedData", []) if f["name"] == "Rarity"), None),
            "image": (prod.get("imageUrl") or "").replace("_200w.", "_400w.") if prod.get("imageCount") else None,
            "url": prod.get("url"),
            "product_id": prod["productId"],
            "variant": variant,
        }

    cards = {}
    for prod, raw_number, _ in plain:
        found = subtypes.get(prod["productId"], set())
        variants = [SUBTYPES[s] for s in found if s in SUBTYPES] or ["normal"]  # unpriced: assume plain print
        for variant in variants:
            cards.setdefault((key_of(raw_number), variant), card(prod, raw_number, variant))
    for prod, raw_number, _, variant in patterned:
        cards.setdefault((key_of(raw_number), variant), card(prod, raw_number, variant))

    # Deck-exclusive prints (e.g. a non-holo of a holo rare from a pre-built deck) live in
    # TCGplayer's "Deck Exclusives" group; they belong to this set when the printed total
    # is the set's official count and the number + name match one of its cards.
    if deck_group_id and official and not reprint:
        names = {(c["key"], fold(c["name"])) for c in cards.values()}
        deck_prices = defaultdict(set)
        for p in get_json(f"{TCGCSV}/{deck_group_id}/prices", f"tcgcsv_{deck_group_id}_prices", refresh)["results"]:
            deck_prices[p["productId"]].add(p["subTypeName"])
        for prod in get_json(f"{TCGCSV}/{deck_group_id}/products", f"tcgcsv_{deck_group_id}_products",
                             refresh)["results"]:
            raw_number = product_number(prod) or ""
            total = raw_number.split("/")[1].strip() if "/" in raw_number else ""
            if not (total.isdigit() and int(total) == official
                    and (number_key(raw_number), fold(clean_card_name(prod["name"]))) in names):
                continue
            found = deck_prices.get(prod["productId"], set())
            for variant in [SUBTYPES[s] for s in found if s in SUBTYPES] or ["normal"]:
                if (number_key(raw_number), variant) not in cards:
                    cards[(number_key(raw_number), variant)] = dict(card(prod, raw_number, variant), deck=True)

    # TCGplayer sometimes lists an image the CDN refuses to serve (403), so check.
    ok = check_images({c["image"] for c in cards.values() if c["image"]}, refresh)
    for c in cards.values():
        if c["image"] and not ok[c["image"]]:
            c["image"] = None

    # Products without a working picture (often new pattern reverses) borrow the
    # image of another print of the same number; the variant badge tells them apart.
    # Every card also gets that plain image as a browser-side fallback.
    by_number = defaultdict(list)
    for (key, variant), card in cards.items():
        if card["image"]:
            by_number[key].append((list(VARIANTS).index(variant), card["image"]))
    for (key, _), card in cards.items():
        plain = min(by_number[key])[1] if by_number[key] else None
        if not card["image"] and plain:
            card["image"] = plain
            card["image_borrowed"] = True
        if plain and plain != card["image"]:
            card["image_fallback"] = plain

    return cards, dict(unrecognised), reprint


def write_stub(slug, name):
    PAGES_DIR.mkdir(exist_ok=True)
    (PAGES_DIR / f"{slug}.md").write_text(
        "---\n"
        "# Generated by pokemon/build_catalogue.py; do not edit by hand.\n"
        "generated: true\n"
        "layout: pokemon-set\n"
        f"title: {json.dumps(name, ensure_ascii=False)}\n"
        f"set_slug: {json.dumps(slug)}\n"
        f"permalink: /pokemon-tcg/{slug}/\n"
        "---\n"
    )


def remove_stale(slugs):
    for path in PAGES_DIR.glob("*.md"):
        if path.stem not in slugs and "\ngenerated: true\n" in path.read_text():
            path.unlink()
            print(f"  removed stale page {path.relative_to(ROOT)}")
    for path in DATA_DIR.glob("*.json"):
        if path.stem != "sets" and path.stem not in slugs:
            path.unlink()
            print(f"  removed stale data {path.relative_to(ROOT)}")


class Universe:
    """Everything a script needs to know about one set's complete card list."""
    def __init__(self, set_name, tset, detail, group, release_date, series, reprint, cards, ov):
        self.set_name, self.tset, self.detail, self.group = set_name, tset, detail, group
        self.release_date, self.series, self.reprint, self.cards, self.ov = release_date, series, reprint, cards, ov


def set_universe(set_name, cats, overrides, deck_group, refresh, lists=None, log=print):
    """The complete card list of a set: booster cards with their variants, subsets, deck
    exclusives and special prints (stamps, promos, ...), keyed (card key, variant).
    Returns (Universe, None) or (None, problem). Used for the public site and master sets
    alike, so both always list the same cards."""
    ov = overrides.get(set_name) or {}
    tset, how = cats.match_tcgdex(set_name, ov.get("tcgdex"))
    if not tset and cats.reload_once():
        tset, how = cats.match_tcgdex(set_name, ov.get("tcgdex"))
    detail = cats.tcgdex_detail(tset["id"]) if tset else {}
    if tset and how != "exact":
        log(f"  TCGdex: {how} match -> {tset['name']} ({tset['id']})")

    fallback = (lists or {}).get(set_name, {})
    release_date = detail.get("releaseDate") or fallback.get("release_date")
    series = (detail.get("serie") or {}).get("name") or fallback.get("series") or "Other"

    group, how = cats.match_group(set_name, release_date, ov.get("tcgplayer"))
    if not group and cats.reload_once():
        group, how = cats.match_group(set_name, release_date, ov.get("tcgplayer"))
    if not group:
        return None, (f"{set_name}: no TCGplayer group found; "
                      f"add a 'tcgplayer: <groupId>' entry to {OVERRIDES.relative_to(ROOT)}")
    if how != "exact":
        log(f"  TCGplayer: {how} match -> {group['name']} ({group['groupId']})")

    extra = cats.subset_groups(group, ov.get("tcgplayer_extra"))
    for g in extra:
        log(f"  TCGplayer: + subset {g['name']} ({g['groupId']})")
    official = (detail.get("cardCount") or {}).get("official")
    cards, unrecognised, reprint = build_card_list(
        [group["groupId"]] + [g["groupId"] for g in extra], refresh, deck_group, official)
    for label, n in sorted(unrecognised.items()):
        log(f"  note: skipped {n} product(s) labelled '({label})' (duplicate of a plain card number)")
    decks = sum(1 for c in cards.values() if c.get("deck"))
    if decks:
        log(f"  + {decks} deck-exclusive print(s)")
    if reprint:
        log("  cards carry their original numbers (e.g. 114/264): matched on the full number")
    else:
        specials, notes = special_prints(set_name, official, cards, cats.groups, refresh)
        for k, c in specials.items():
            cards.setdefault(k, c)
        if specials:
            ok = check_images({c["image"] for c in specials.values() if c["image"]}, refresh)
            for c in specials.values():
                if c["image"] and not ok[c["image"]]:
                    c["image"] = None
            # e.g. staff promos often have no picture: borrow the same card's other print
            pictures = {}
            for c in cards.values():
                if c.get("image"):
                    pictures.setdefault(c["key"], c["image"])
            for c in specials.values():
                if not c["image"] and c["key"] in pictures:
                    c["image"], c["image_borrowed"] = pictures[c["key"]], True
            kinds = defaultdict(int)
            for _, kind in specials:
                kinds[VARIANTS[kind]] += 1
            log(f"  + special prints: {', '.join(f'{n} {k}' for k, n in sorted(kinds.items()))}")
        for note in notes:
            log(f"  note: skipped {note}")
    return Universe(set_name, tset, detail, group, release_date, series, reprint, cards, ov), None


def resolve_entries(universe, entries, where_of=None):
    """Spreadsheet rows -> card keys. entries: {(full_key(number), variant): qty}.
    '25' and '025/217' both mean card 25 in a normal set; in a reprint-numbered set '114'
    works when only one card is numbered 114. Returns ({(key, variant): qty}, problems,
    {(key, variant): [row locations]}). where_of(key, variant) names the rows for messages."""
    where_of = where_of or (lambda key, variant: "")
    cards, set_name = universe.cards, universe.set_name
    numbers = {k for k, _ in cards}
    resolved, where, problems = defaultdict(int), defaultdict(list), []
    for (key, variant), qty in entries.items():
        loc = where_of(key, variant)
        prefix = f"{loc}: " if loc else ""
        bare = key.split("/")[0]
        if key in numbers:
            matches = [key]
        elif bare in numbers:
            matches = [bare]
        else:
            matches = sorted({k for k in numbers if k.split("/")[0] == bare}, key=number_sort)
        if len(matches) != 1:
            if matches:
                problems.append(f"{prefix}{set_name} #{key}: several cards have this number "
                                f"({', '.join(matches)}); type the full printed number")
            else:
                problems.append(f"{prefix}{set_name} #{key} ({variant}): card number not found in this set")
            continue
        if (matches[0], variant) not in cards:
            have = sorted({v for k, v in cards if k == matches[0]}, key=list(VARIANTS).index)
            problems.append(f"{prefix}{set_name} #{key}: no '{variant}' variant (known: {', '.join(have)})")
            continue
        resolved[(matches[0], variant)] += qty
        if loc:
            where[(matches[0], variant)].append(loc)
    return dict(resolved), problems, where


def ordered_cards(universe, log=print):
    """The universe's cards as a sorted list (copies), with variant labels; keeps 'key'."""
    order = list(VARIANTS)
    card_list = sorted((dict(c) for c in universe.cards.values()),
                       key=lambda c: (number_sort(c["key"]), order.index(c["variant"])))
    if universe.reprint:  # list them in TCGdex's set order (001, 002, ...) rather than by original number
        by_name = match_by_name(card_list, universe.detail.get("cards", []))
        for c in card_list:
            if (c["number"], c["name"]) in by_name:
                c["tcgdex_local"] = by_name[(c["number"], c["name"])]["localId"]
        missing = sorted({c["name"] for c in card_list if "tcgdex_local" not in c})
        if missing:
            log(f"  note: not on TCGdex, listed last: {', '.join(missing)}")
        card_list.sort(key=lambda c: ("tcgdex_local" not in c, number_sort(number_key(c.get("tcgdex_local", "0"))),
                                      number_sort(c["key"]), order.index(c["variant"])))
    if universe.ov.get("order"):  # explicit order (e.g. the ETB booklet's), by printed number; unlisted last
        keys = {c["key"] for c in card_list}

        def to_key(n):  # like the xlsx: '114' is enough when only one card has that number
            k = key_of_number(n, universe.reprint)
            same = [x for x in keys if x.split("/")[0] == k.split("/")[0]]
            return k if k in keys or len(same) != 1 else same[0]
        wanted = [to_key(n) for n in universe.ov["order"]]
        unknown = [n for n, k in zip(universe.ov["order"], wanted) if k not in keys]
        if unknown:
            log(f"  note: 'order' in {OVERRIDES.name} lists numbers not in this set: {', '.join(map(str, unknown))}")
        position = {k: i for i, k in reversed(list(enumerate(wanted)))}
        card_list.sort(key=lambda c: position.get(c["key"], len(position)))  # stable: keeps variant order
    for c in card_list:
        c["variant_label"] = VARIANTS[c["variant"]] + (" (deck)" if c.get("deck") else "")
        if c.get("special") and fold(c["special"]) != fold(VARIANTS[c["variant"]]):
            c["variant_label"] += f" ({c['special']})"  # e.g. "Stamped (Pitch Black Stamped)"
    return card_list


def deck_group_id(cats):
    return next((g["groupId"] for g in cats.groups if fold(g["name"]) == "deck exclusives"), None)


def main():
    parser = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    parser.add_argument("--refresh", action="store_true", help="ignore the API cache and refetch everything")
    args = parser.parse_args()

    owned, rows_of, lists, report = read_catalogue()
    overrides = read_overrides()
    cats = Catalogues(args.refresh)
    deck_group = deck_group_id(cats)

    by_set = defaultdict(dict)
    for (set_name, key, variant), qty in owned.items():
        by_set[set_name][(key, variant)] = qty

    DATA_DIR.mkdir(parents=True, exist_ok=True)
    summaries = []
    for set_name in sorted(by_set):
        print(f"• {set_name}")
        universe, problem = set_universe(set_name, cats, overrides, deck_group, args.refresh, lists)
        if not universe:
            report.append(f"{problem} ({sum(by_set[set_name].values())} cards skipped)")
            continue
        resolved, problems, _ = resolve_entries(
            universe, by_set[set_name], lambda key, variant: ", ".join(rows_of[(set_name, key, variant)]))
        report += [f"{line}; skipped" for line in problems]

        card_list = ordered_cards(universe)
        for c in card_list:
            c["qty"] = resolved.get((c["key"], c["variant"]), 0)
            del c["key"]

        ov, detail, tset, group = universe.ov, universe.detail, universe.tset, universe.group
        release_date, reprint = universe.release_date, universe.reprint
        slug = slugify(set_name)
        # first logo that loads: override, TCGdex, pokemontcg.io (TCGdex lacks some, and some of its URLs 404)
        logo = None
        for source, url in (("override", lambda: ov.get("logo")),
                            ("TCGdex", lambda: f"{detail['logo']}.png" if detail.get("logo") else None),
                            ("pokemontcg.io", lambda: cats.fallback_logo(set_name, release_date))):
            candidate = url()
            if not candidate:
                continue
            if check_images([candidate], False)[candidate]:
                logo = candidate
                if source == "pokemontcg.io":
                    print("  logo: from pokemontcg.io")
                break
            print(f"  logo: {source} logo {candidate} does not load")
        if not logo:
            print("  logo: none found; add a 'logo: <url>' override, the set name is shown as text meanwhile")
        summary = {
            "slug": slug,
            "name": set_name,
            "series": universe.series,
            "release_date": release_date,
            "logo": logo,
            "tcgdex_id": tset["id"] if tset else None,
            "tcgplayer_group_id": group["groupId"],
            "tcgplayer_group_name": group["name"],
            "owned_count": sum(1 for c in card_list if c["qty"] > 0),
            "total_count": len(card_list),
            "owned_qty": sum(c["qty"] for c in card_list),
            "reprint_numbers": reprint,  # numbered as the original prints (Classic Collection)
        }
        summaries.append(summary)
        (DATA_DIR / f"{slug}.json").write_text(
            json.dumps(dict(summary, cards=card_list), ensure_ascii=False, indent=1) + "\n", encoding="utf_8")
        write_stub(slug, set_name)

    # Newest first; Liquid's group_by keeps this order, so series are ordered by their newest set.
    summaries.sort(key=lambda s: (s["release_date"] or "", s["name"]), reverse=True)
    (DATA_DIR / "sets.json").write_text(json.dumps(summaries, ensure_ascii=False, indent=1) + "\n", encoding="utf_8")
    remove_stale({s["slug"] for s in summaries})

    print(f"\nWrote {len(summaries)} set(s) to {DATA_DIR.relative_to(ROOT)}/")
    if report:
        print(f"\n{len(report)} problem(s) — these xlsx rows are NOT on the site:")
        for line in report:
            print(f"  - {line}")
        sys.exit(1)
    print("All xlsx rows matched.")


if __name__ == "__main__":
    main()
