#!/usr/bin/env python3
# Run from the repo root (the folder containing pokemon/); close the xlsx files in Excel first.
#
# Windows (PowerShell / VS Code terminal):
#   pokemon\.venv\Scripts\python.exe pokemon\add_entries.py --to public --set ASC --dry-run
#   pokemon\.venv\Scripts\python.exe pokemon\add_entries.py --to public --set ASC
#   pokemon\.venv\Scripts\python.exe pokemon\build_catalogue.py      # then rebuild the site data
#   pokemon\.venv\Scripts\python.exe pokemon\add_entries.py --to masterset:PBL
#   pokemon\.venv\Scripts\python.exe pokemon\add_entries.py --to masterset:Kingdra --set "Paradox Rift"
# macOS:
#   pokemon/.venv/bin/python pokemon/add_entries.py --to public --set ASC
#   pokemon/.venv/bin/python pokemon/build_catalogue.py
#
# Replace ASC with the set's abbreviation (TEF, PRE, ...). First time on a machine:
# see the venv setup at the top of build_catalogue.py.
"""Add the cards in pokemon/source/new_entries.xlsx to the public collection or a master set.

--to is required and has no default:
  --to public               the public bulk catalogue (pokemon_bulk_catalogue.xlsx), on the
                            tab of the set given with --set (abbreviation, e.g. ASC, or name)
  --to masterset:<name>     a master set in the private pokemon_mastersets.xlsx (local only),
                            e.g. masterset:PBL or "masterset:Pitch Black"; Pokémon / illustrator
                            master sets (masterset:Kingdra) also need --set for the cards' set

Every row (number, variant, qty) is first checked against the destination's card list
(the set's complete list, or the master set's rows). If any row doesn't match, all of them
are listed and nothing is written. Then:

  - same number + variant already there -> its qty is increased by the new qty
                                            (--replace: set to the new qty); a master set
                                            card you already had is flagged as a duplicate
  - public, not on the tab yet          -> appended as a new row

A set without a public tab gets one, copied from the Template tab. The workbook is backed up
to pokemon/.cache/backups/ first, and new_entries.xlsx is emptied afterwards so the same
cards can't be added twice (--keep leaves it as is).

Usage:  python pokemon/add_entries.py --to (public | masterset:<name>) [--set ABBR]
                                      [--replace] [--keep] [--dry-run]
"""

import argparse
import shutil
import sys
from concurrent.futures import ThreadPoolExecutor
from copy import copy
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import build_catalogue as bc  # noqa: E402  (also silences openpyxl/urllib3 warnings)

import openpyxl  # noqa: E402
from openpyxl.styles import Font, PatternFill  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

NEW_ENTRIES = bc.SOURCE_DIR / "new_entries.xlsx"
BACKUP_DIR = bc.CACHE_DIR / "backups"
HEADERS = ("number", "variant", "qty")
FIRST_ROW = 4  # catalogue tabs: set_name in B1, headers in row 3, cards from row 4


def rel(path):
    """Path as shown to the user: relative to the repo when inside it."""
    try:
        return path.relative_to(bc.ROOT)
    except ValueError:
        return path


def fail(message):
    print(f"error: {message}", file=sys.stderr)
    sys.exit(1)


def clean_number(raw, keep_total=False):
    """'025/217' -> '25', 25.0 -> '25', 'tg05' -> 'TG05' (text as typed otherwise).
    keep_total: '106/105' -> '106/105', for sets numbered as the original prints
    (Classic Collection), where only the set total tells some cards apart."""
    if isinstance(raw, float) and raw.is_integer():
        raw = int(raw)
    text, _, total = str(raw).strip().partition("/")
    clean = lambda t: str(int(t)) if t.isdigit() else t
    text, total = clean(text.strip().upper()), clean(total.strip().upper())
    return f"{text}/{total}" if keep_total and total else text


def entry_key(number, shared=None):
    """Matching key. shared: None for a normal set, else (a set numbered as the original
    prints) the numbers several of its cards have, which need their set total: '114' and
    '114/264' are the same card, '106/105' and '106/160' are not."""
    if shared is None:
        return bc.number_key(number)
    return bc.full_key(number) if bc.number_key(number) in shared else bc.number_key(number)


def shared_numbers(set_name):
    """None unless the last build found this set numbered as the original prints; then
    the card numbers that more than one of its cards has (e.g. {'106', '11'})."""
    import json
    sets_file = bc.DATA_DIR / "sets.json"
    if not sets_file.exists():
        return None
    summary = next((s for s in json.loads(sets_file.read_text(encoding="utf_8")) if s["name"] == set_name), None)
    if not summary or not summary.get("reprint_numbers"):
        return None
    cards = json.loads((bc.DATA_DIR / f"{summary['slug']}.json").read_text(encoding="utf_8"))["cards"]
    full = {bc.full_key(c["number"]) for c in cards}
    bare = [k.split("/")[0] for k in full]
    return {n for n in bare if bare.count(n) > 1}


def ensure_closed(*paths):
    """Excel on macOS doesn't lock files, so a save here would be lost (or
    clobbered) when Excel saves its open copy. Its ~$ owner file gives it away."""
    for path in paths:
        if (path.parent / f"~${path.name}").exists():
            fail(f"{path.name} is open in Excel; close it first")


def create_new_entries_file():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "New entries"
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    for col, width in zip("ABC", (12, 16, 8)):
        ws.column_dimensions[col].width = width
    fill = PatternFill("solid", fgColor="FFF9DB")  # same yellow as the catalogue
    for row in ws.iter_rows(min_row=2, max_row=500, max_col=3):
        for cell in row:
            cell.fill = fill
    variants = DataValidation(type="list", formula1='"' + ",".join(bc.VARIANTS) + '"', allow_blank=True)
    variants.add("B2:B500")
    qty = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="1", allow_blank=True)
    qty.add("C2:C500")
    ws.add_data_validation(variants)
    ws.add_data_validation(qty)
    wb.save(NEW_ENTRIES)


def read_new_entries(shared=None):
    """{(number_key, variant): (number, variant, qty)} plus the row numbers read."""
    wb = openpyxl.load_workbook(NEW_ENTRIES, data_only=True)
    ws = wb.worksheets[0]
    header = [str(c.value).strip().lower() if c.value is not None else "" for c in ws[1]]
    if not all(h in header for h in HEADERS):
        fail(f"{NEW_ENTRIES.name}: row 1 must be the headers number / variant / qty")
    col = {h: header.index(h) for h in HEADERS}

    entries, rows, problems = {}, [], []
    for rownum, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        number, variant, qty = (row[col[h]] if col[h] < len(row) else None for h in HEADERS)
        if all(v in (None, "") for v in (number, variant, qty)):
            continue
        rows.append(rownum)
        variant = str(variant or "").strip().lower()
        if number in (None, ""):
            problems.append(f"row {rownum}: missing number")
            continue
        if variant not in bc.VARIANTS:
            problems.append(f"row {rownum}: #{number}: unknown variant '{variant}'")
            continue
        try:
            qty = int(qty) if qty not in (None, "") else 1
        except (TypeError, ValueError):
            problems.append(f"row {rownum}: #{number}: qty '{qty}' is not a number")
            continue
        if qty <= 0:
            problems.append(f"row {rownum}: #{number}: qty must be 1 or more")
            continue
        number = clean_number(number, shared is not None)
        key = (entry_key(number, shared), variant)
        if key in entries:  # same card twice in new_entries: add them up
            qty += entries[key][2]
        entries[key] = (number, variant, qty)

    if problems:
        soft_exit(f"{NEW_ENTRIES.name} has {len(problems)} problem(s); nothing was written:\n  " + "\n  ".join(problems))
    return entries, rows


# ---------------------------------------------------------------- finding the tab

class SetLookup:
    """Official abbreviations per set name: TCGdex first, TCGplayer's as a fallback."""

    def __init__(self):
        self.cats = bc.Catalogues(False)
        self.overrides = bc.read_overrides()

    def abbreviations(self, set_name):
        ov = self.overrides.get(set_name) or {}
        found = set()
        tset, _ = self.cats.match_tcgdex(set_name, ov.get("tcgdex"))
        release_date = None
        if tset:
            detail = self.cats.tcgdex_detail(tset["id"])
            release_date = detail.get("releaseDate")
            abbr = (detail.get("abbreviation") or {}).get("official")
            if abbr:
                found.add(abbr.upper())
        group, _ = self.cats.match_group(set_name, release_date, ov.get("tcgplayer"))
        if group and group.get("abbreviation"):
            found.add(group["abbreviation"].upper())
        return found


def set_tabs(wb):
    """{tab title: set_name in B1} for every set tab that has a set name."""
    tabs = {}
    for ws in wb.worksheets:
        if ws.title in bc.NON_COLLECTION_SHEETS or ws.title == "Template":
            continue
        name = ws["B1"].value
        if name not in (None, ""):
            tabs[ws.title] = str(name).strip()
    return tabs


def find_set(wb, wanted, lookup):
    """(tab title or None, set_name). Matches an existing tab first, then any set
    in the Lists sheet (whose tab will be created)."""
    want = wanted.strip().upper()
    tabs = set_tabs(wb)

    # an exact set or tab name also works, e.g. --set "Ascended Heroes"
    for title, name in tabs.items():
        if want in (title.upper(), name.upper()):
            return title, name

    hits = [(title, name) for title, name in tabs.items() if want in lookup.abbreviations(name)]
    if len(hits) > 1:
        fail(f"'{wanted}' matches several tabs ({', '.join(t for t, _ in hits)}); "
             f"use the full set name, e.g. --set \"{hits[0][1]}\"")
    if hits:
        return hits[0]

    # no tab yet: look through every set in the Lists sheet (dropdown source)
    names = [str(r[0]).strip() for r in wb["Lists"].iter_rows(min_row=2, max_col=1, values_only=True) if r[0]]
    exact = [n for n in names if n.upper() == want]
    if exact:
        return None, exact[0]
    print("Looking up set abbreviations (slow only the first time)...")
    with ThreadPoolExecutor(max_workers=8) as pool:
        abbrs = list(pool.map(lookup.abbreviations, names))
    hits = [n for n, a in zip(names, abbrs) if want in a]
    if len(hits) > 1:
        fail(f"'{wanted}' matches several sets ({', '.join(hits)}); use the full set name instead")
    if not hits:
        known = ", ".join(f"{sorted(lookup.abbreviations(n)) or ['?']} {n}" for n in tabs.values())
        fail(f"no set with abbreviation '{wanted}'. Sets with a tab: {known or 'none'}")
    return None, hits[0]


def create_tab(wb, set_name):
    template = wb["Template"]
    ws = wb.copy_worksheet(template)  # copies cells, styles and column widths only
    title = "".join(ch for ch in set_name if ch not in "[]:*?/\\")[:31]
    if title in wb.sheetnames:
        title = title[:28] + " 2"
    ws.title = title
    ws["B1"] = set_name
    ws.freeze_panes = template.freeze_panes
    for dv in template.data_validations.dataValidation:
        new = copy(dv)
        new.sqref = copy(dv.sqref)
        ws.add_data_validation(new)
    wb.move_sheet(ws, offset=wb.index(wb["Lists"]) - wb.index(ws))  # keep Lists last
    return ws


# ---------------------------------------------------------------- merging

def merge(ws, entries, replace, shared=None):
    existing = {}  # (number_key, variant) -> row
    last_row = FIRST_ROW - 1
    for row in ws.iter_rows(min_row=FIRST_ROW, max_col=3):
        number, variant, qty = (c.value for c in row)
        if all(v in (None, "") for v in (number, variant, qty)):
            continue
        last_row = row[0].row
        if number not in (None, "") and variant:
            key = (entry_key(clean_number(number, shared is not None), shared), str(variant).strip().lower())
            existing.setdefault(key, row[0].row)  # duplicates: update the first one

    changes = []
    for key, (number, variant, qty) in sorted(entries.items(), key=lambda kv: bc.number_sort(kv[0][0])):
        if key in existing:
            cell = ws.cell(existing[key], 3)
            try:
                old = int(cell.value) if cell.value not in (None, "") else 1  # blank counts as 1
            except (TypeError, ValueError):
                old = 0
            new = qty if replace else old + qty
            cell.value = new
            changes.append(f"  updated  row {existing[key]:>4}: #{number} {variant}  qty {old} -> {new}")
        else:
            last_row += 1
            ws.cell(last_row, 1, number)
            ws.cell(last_row, 2, variant)
            ws.cell(last_row, 3, qty)
            changes.append(f"  added    row {last_row:>4}: #{number} {variant}  qty {qty}")
    return changes


def clear_entries(rows):
    wb = openpyxl.load_workbook(NEW_ENTRIES)
    ws = wb.worksheets[0]
    for rownum in rows:
        for col in range(1, 4):
            ws.cell(rownum, col).value = None
    wb.save(NEW_ENTRIES)


def soft_exit(message):
    """Stop without changing anything (exit status 0: nothing went wrong, there's just nothing to do)."""
    print(message)
    sys.exit(0)


def validation_entries():
    """new_entries.xlsx rows keyed for validation: {(full_key(number), variant): qty}, keeping
    printed totals ('106/105') so reprint-numbered sets can be told apart."""
    entries, _ = read_new_entries(shared=set())
    return {(bc.full_key(number), variant): qty for number, variant, qty in entries.values()}


def check_entries(universe, entries, allowed=None, where=""):
    """Problems with entries against a set's complete card list and, for a master set, the
    (key, variant) pairs it lists (allowed). Returns (resolved, problems)."""
    resolved, problems, _ = bc.resolve_entries(universe, entries)
    if allowed is not None:
        for key, variant in resolved:
            if (key, variant) not in allowed:
                problems.append(f"{universe.set_name} #{key} ({variant}): not part of {where}")
    return resolved, problems


def report_problems(problems, destination):
    soft_exit(f"{NEW_ENTRIES.name} has {len(problems)} entr{'y that doesn' if len(problems) == 1 else 'ies that don'}'t "
              f"match {destination}; nothing was written:\n  " + "\n  ".join(problems))


def add_to_public(args, catalogue_args):
    if not args.set:
        soft_exit("--to public needs --set (the set's abbreviation or name, e.g. --set TEF); nothing was written.")
    entries, rows = read_new_entries()
    wb = openpyxl.load_workbook(bc.XLSX)
    title, set_name = find_set(wb, args.set, SetLookup())

    universe, problem = bc.set_universe(set_name, *catalogue_args, log=lambda *a: None)
    if not universe:
        soft_exit(f"{problem}; nothing was written.")
    _, problems = check_entries(universe, validation_entries())
    if problems:
        report_problems(problems, f"the card list of {set_name}")

    if title is None:
        ws = create_tab(wb, set_name)
        print(f"{set_name}: no tab yet, created '{ws.title}' from the Template tab")
    else:
        ws = wb[title]
    print(f"Public collection, {set_name} (tab '{ws.title}'), "
          f"{len(entries)} entr{'y' if len(entries) == 1 else 'ies'}:")
    shared = shared_numbers(set_name)
    if shared is not None:  # keep '106/105' whole: several cards are numbered 106
        entries, rows = read_new_entries(shared)
        print(f"(numbered as the original prints; type the full number for #{', #'.join(sorted(shared))}, "
              f"e.g. 106/105)")
    for line in merge(ws, entries, args.replace, shared):
        print(line)
    return wb, bc.XLSX, rows


def add_to_masterset(args, name):
    try:
        import masterset_lib as ml
    except ImportError:
        soft_exit("The master set scripts aren't on this machine (they're local only); nothing was written.")
    ensure_closed(ml.XLSX)
    if not ml.XLSX.exists():
        soft_exit(f"{rel(ml.XLSX)} doesn't exist yet; start a master set with masterset_setup.py first.")
    ctx = ml.Context(quiet=True)
    wb = ml.open_workbook()
    tab, ws, legacy = ml.find_tab(wb, name)
    if not tab:  # maybe a set abbreviation (PBL) of a set master set
        for set_name in ctx.set_names(name):
            tab, ws, legacy = ml.find_tab(wb, set_name)
            if tab:
                break
    if not tab:
        known = ", ".join(sorted(ml.masterset_tabs(wb))) or "none"
        soft_exit(f"No master set '{name}' in {rel(ml.XLSX)} (master sets: {known}); nothing was written.")
    if legacy:
        soft_exit(f"Master set '{tab}' still has the old layout; it has to be converted first. Nothing was written.")

    rows = ml.read_rows(ws)
    sets_in_tab = list(dict.fromkeys(r["set"] for r in rows))
    if ml.tab_kind(ws) == "set" and not args.set:
        set_name = sets_in_tab[0] if sets_in_tab else tab
    else:
        if not args.set:
            example = sets_in_tab[0] if sets_in_tab else "Pitch Black"
            soft_exit(f"Master set '{tab}' spans {len(sets_in_tab)} sets: say which one the cards are from with "
                      f"--set (e.g. --set \"{example}\"); nothing was written.")
        candidates = [s for s in ctx.set_names(args.set) if s in sets_in_tab]
        if len(candidates) != 1:
            soft_exit(f"Set '{args.set}' isn't part of master set '{tab}' (its sets: {', '.join(sets_in_tab)}); "
                      f"nothing was written.")
        set_name = candidates[0]

    universe = ctx.universe(set_name)
    if not universe:
        soft_exit(f"No card list for {set_name}; nothing was written.")
    by_card = {}  # (key, variant) -> master set row
    for r in rows:
        if r["set"] == set_name:
            resolved, _, _ = bc.resolve_entries(universe, {(bc.full_key(r["number"]), r["variant"]): 1})
            for k in resolved:
                by_card[k] = r
    resolved, problems = check_entries(universe, validation_entries(), set(by_card), f"master set '{tab}'")
    if problems:
        report_problems(problems, f"master set '{tab}' ({set_name})")

    _, entry_rows = read_new_entries(shared=set())
    print(f"Master set '{tab}', {set_name}, {len(resolved)} entr{'y' if len(resolved) == 1 else 'ies'}:")
    for (key, variant), qty in sorted(resolved.items(), key=lambda kv: bc.number_sort(kv[0][0])):
        r = by_card[(key, variant)]
        cell = ws.cell(r["row"], ml.HEADERS.index("qty") + 1)
        old = r["qty"]
        new = qty if args.replace else old + qty
        cell.value = new
        note = f"   <- already had {old}: this is a duplicate" if old and not args.replace else ""
        print(f"  row {r['row']:>4}: #{r['number']} {r['name']} ({variant})  qty {old} -> {new}{note}")
    args.masterset_tab = tab  # refreshed (PDF + CardTrader wishlist) after saving
    return wb, ml.XLSX, entry_rows


def main():
    parser = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    parser.add_argument("--to", metavar="DESTINATION",
                        help="where the cards go: 'public' (bulk collection) or 'masterset:<name>' "
                             "(e.g. masterset:PBL, \"masterset:Pitch Black\", masterset:Kingdra); required")
    parser.add_argument("--set", metavar="ABBR",
                        help="official set abbreviation, e.g. ASC or TEF (a full set name also works); "
                             "required for --to public and for Pokémon / illustrator master sets")
    parser.add_argument("--replace", action="store_true",
                        help="set the qty of existing rows to the new qty instead of adding to it")
    parser.add_argument("--keep", action="store_true", help="don't empty new_entries.xlsx afterwards")
    parser.add_argument("--dry-run", action="store_true", help="show what would change, write nothing")
    args = parser.parse_args()

    if not args.to:
        soft_exit("A destination is required: --to public, or --to masterset:<name> (e.g. masterset:PBL). "
                  "Nothing was written.")
    destination = args.to.strip()
    if destination.lower() != "public" and not destination.lower().startswith("masterset:"):
        soft_exit(f"Unknown destination '{destination}': use --to public or --to masterset:<name>. "
                  "Nothing was written.")

    if not NEW_ENTRIES.exists():
        create_new_entries_file()
        print(f"Created {rel(NEW_ENTRIES)}; fill in number / variant / qty and run again.")
        return
    ensure_closed(bc.XLSX, NEW_ENTRIES)
    entries, _ = read_new_entries()
    if not entries:
        print(f"{NEW_ENTRIES.name} is empty; nothing to add.")
        return

    if destination.lower() == "public":
        cats = bc.Catalogues(False)
        wb, path, rows = add_to_public(args, (cats, bc.read_overrides(), bc.deck_group_id(cats), False))
    else:
        name = destination.split(":", 1)[1].strip()
        if not name:
            soft_exit("Name the master set: --to masterset:<name>. Nothing was written.")
        wb, path, rows = add_to_masterset(args, name)

    if args.dry_run:
        print("\nDry run: nothing was written.")
        return

    BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    backup = BACKUP_DIR / f"{path.stem}-{datetime.now():%Y%m%d-%H%M%S}.xlsx"
    shutil.copy2(path, backup)
    wb.save(path)
    print(f"\nSaved {path.name} (previous version: {rel(backup)})")
    if not args.keep:
        clear_entries(rows)
        print(f"Emptied {NEW_ENTRIES.name}.")
    if path == bc.XLSX:
        print("Next: pokemon/.venv/bin/python pokemon/build_catalogue.py")
    else:  # a master set changed: new PDF + fair page, CardTrader wishlist replaced
        import masterset_placeholders
        print()
        print(f"Refreshing master set '{args.masterset_tab}'...")
        for line in masterset_placeholders.refresh(args.masterset_tab):
            print(line)


if __name__ == "__main__":
    main()
